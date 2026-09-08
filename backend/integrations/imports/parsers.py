import csv
import hashlib
import io
from datetime import date, datetime
from pathlib import Path

from rest_framework import serializers


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 80


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _validate_table(headers, rows):
    if not headers or not any(str(item).strip() for item in headers):
        raise serializers.ValidationError({"file": "The import file has no header row."})
    if len(headers) > MAX_COLUMNS:
        raise serializers.ValidationError(
            {"file": f"Import files can contain at most {MAX_COLUMNS} columns."}
        )
    if len(rows) > MAX_ROWS:
        raise serializers.ValidationError(
            {"file": f"Import files can contain at most {MAX_ROWS} data rows."}
        )
    cleaned_headers = [_cell_text(item) for item in headers]
    if any(not item for item in cleaned_headers):
        raise serializers.ValidationError(
            {"file": "Every import column must have a non-empty header."}
        )
    return cleaned_headers


def _parse_csv(raw):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise serializers.ValidationError(
            {"file": "CSV files must be saved as UTF-8."}
        ) from exc

    try:
        reader = csv.reader(io.StringIO(text, newline=""))
        records = []
        for index, record in enumerate(reader):
            if len(record) > MAX_COLUMNS:
                raise serializers.ValidationError(
                    {"file": f"Import files can contain at most {MAX_COLUMNS} columns."}
                )
            records.append(record)
            if index > MAX_ROWS:
                raise serializers.ValidationError(
                    {"file": f"Import files can contain at most {MAX_ROWS} data rows."}
                )
    except csv.Error as exc:
        raise serializers.ValidationError({"file": "The CSV file could not be parsed."}) from exc

    if not records:
        raise serializers.ValidationError({"file": "The CSV file is empty."})
    headers = records[0]
    width = len(headers)
    rows = []
    for record in records[1:]:
        padded = list(record[:width]) + [""] * max(0, width - len(record))
        values = [_cell_text(item) for item in padded]
        if any(values):
            rows.append(values)
    return _validate_table(headers, rows), rows


def _parse_xlsx(raw):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise serializers.ValidationError(
            {"file": "Excel import support is not installed on the server."}
        ) from exc

    try:
        workbook = load_workbook(
            filename=io.BytesIO(raw),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise serializers.ValidationError(
            {"file": "The Excel workbook could not be opened."}
        ) from exc

    try:
        worksheet = workbook.active
        if worksheet.max_column > MAX_COLUMNS:
            raise serializers.ValidationError(
                {"file": f"Import files can contain at most {MAX_COLUMNS} columns."}
            )
        if worksheet.max_row > MAX_ROWS + 1:
            raise serializers.ValidationError(
                {"file": f"Import files can contain at most {MAX_ROWS} data rows."}
            )

        iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = next(iterator)
        except StopIteration as exc:
            raise serializers.ValidationError({"file": "The Excel workbook is empty."}) from exc

        width = len(headers)
        rows = []
        for record in iterator:
            values = [_cell_text(item) for item in list(record)[:width]]
            values += [""] * max(0, width - len(values))
            if any(values):
                rows.append(values)
            if len(rows) > MAX_ROWS:
                raise serializers.ValidationError(
                    {"file": f"Import files can contain at most {MAX_ROWS} data rows."}
                )
        return _validate_table(headers, rows), rows
    finally:
        workbook.close()


def parse_uploaded_table(upload):
    filename = Path(str(getattr(upload, "name", "") or "")).name
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise serializers.ValidationError(
            {"file": "Upload a .csv or .xlsx file. Legacy .xls and macro workbooks are not accepted."}
        )

    declared_size = int(getattr(upload, "size", 0) or 0)
    if declared_size > MAX_UPLOAD_BYTES:
        raise serializers.ValidationError(
            {"file": "Import files cannot exceed 5 MB."}
        )

    raw = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise serializers.ValidationError(
            {"file": "Import files cannot exceed 5 MB."}
        )
    if not raw:
        raise serializers.ValidationError({"file": "The import file is empty."})

    if extension == ".csv":
        headers, rows = _parse_csv(raw)
        file_type = "csv"
    else:
        headers, rows = _parse_xlsx(raw)
        file_type = "xlsx"

    return {
        "filename": filename[:255],
        "fileType": file_type,
        "sha256": hashlib.sha256(raw).hexdigest(),
        "headers": headers,
        "rows": rows,
    }
